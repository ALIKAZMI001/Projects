"""Mark attendance by hand gesture (thumb up = present, thumb down = absent).

Reads student names from ``attendance.xlsx`` (column A) and writes "P"/"A" into
column C, one student at a time.

Run:  python attendance.py      (press 'q' to skip the current student)
"""

import time
from pathlib import Path

import cv2
import mediapipe as mp
from openpyxl import load_workbook

BASE_DIR = Path(__file__).resolve().parent
WORKBOOK_PATH = BASE_DIR / "attendance.xlsx"

START_DELAY_SECONDS = 6
BETWEEN_STUDENTS_SECONDS = 6
GESTURE_MARGIN_PX = 50


def main() -> None:
    cap = cv2.VideoCapture(0)
    if not cap.isOpened():
        raise SystemExit("Could not open webcam.")

    hand_detector = mp.solutions.hands.Hands()
    drawing_utils = mp.solutions.drawing_utils

    workbook = load_workbook(WORKBOOK_PATH)
    sheet = workbook.active
    current_row = 2  # row 1 is the header

    def update_attendance(status: str) -> None:
        sheet.cell(row=current_row, column=3).value = status
        workbook.save(WORKBOOK_PATH)

    print(f"Starting attendance check in {START_DELAY_SECONDS} seconds...")
    time.sleep(START_DELAY_SECONDS)

    while current_row <= sheet.max_row:
        student_name = sheet.cell(row=current_row, column=1).value
        print(f"Checking attendance for {student_name}")

        gesture_detected = False
        while not gesture_detected:
            success, frame = cap.read()
            if not success:
                break

            output = hand_detector.process(cv2.cvtColor(frame, cv2.COLOR_BGR2RGB))

            for hand in output.multi_hand_landmarks or []:
                drawing_utils.draw_landmarks(frame, hand)
                landmarks = hand.landmark

                thumb_tip = landmarks[4]
                wrist = landmarks[0]
                frame_height = frame.shape[0]

                thumb_y = int(thumb_tip.y * frame_height)
                wrist_y = int(wrist.y * frame_height)

                if thumb_y < wrist_y - GESTURE_MARGIN_PX:
                    update_attendance("P")
                    gesture_detected = True
                    print(f"Attendance for {student_name} marked as Present")
                    time.sleep(1)
                elif thumb_y > wrist_y + GESTURE_MARGIN_PX:
                    update_attendance("A")
                    gesture_detected = True
                    print(f"Attendance for {student_name} marked as Absent")
                    time.sleep(1)

            cv2.imshow("Attendance Tracker", frame)
            if cv2.waitKey(1) & 0xFF == ord("q"):
                gesture_detected = True
                break

        current_row += 1
        if current_row <= sheet.max_row:
            print(f"Waiting {BETWEEN_STUDENTS_SECONDS}s before the next student...")
            time.sleep(BETWEEN_STUDENTS_SECONDS)

    cap.release()
    cv2.destroyAllWindows()
    print("Attendance check complete.")


if __name__ == "__main__":
    main()
